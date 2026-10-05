#!/usr/bin/env python3
"""Builds the walkthrough's upload files from the fictional text CVs: Jordan Patel as a PDF, Taylor Reid as a Word
document and a small Northwind logo. Re-run after editing a .txt file:

    python explore/data/make.py
"""
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

HERE = Path(__file__).resolve().parent


def pdf(src: Path, out: Path) -> None:
    page = canvas.Canvas(str(out), pagesize=A4, invariant=1)
    page.setTitle(src.stem.replace("-", " ").title())
    y = A4[1] - 60
    for line in src.read_text(encoding="utf-8").splitlines():
        if y < 60:
            page.showPage()
            y = A4[1] - 60
        page.setFont("Helvetica-Bold" if line and not line.startswith(("-", " ")) and len(line) < 40 else "Helvetica", 10)
        page.drawString(50, y, line)
        y -= 14
    page.save()


def word(src: Path, out: Path) -> None:
    doc = Document()
    lines = src.read_text(encoding="utf-8").splitlines()
    doc.add_heading(lines[0], level=1)
    for line in lines[1:]:
        if line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line and len(line) < 40 and not line.startswith(" "):
            doc.add_heading(line, level=2)
        elif line:
            doc.add_paragraph(line.strip())
    doc.core_properties.author = "HermitShell walkthrough"
    doc.save(str(out))


def logo(out: Path) -> None:
    img = Image.new("RGB", (240, 80), "#0b3d5c")
    draw = ImageDraw.Draw(img)
    draw.ellipse((14, 14, 66, 66), fill="#f2a900")
    draw.text((84, 32), "NORTHWIND", fill="white")
    img.save(out, optimize=True)


if __name__ == "__main__":
    pdf(HERE / "jordan-patel.txt", HERE / "jordan-patel.pdf")
    word(HERE / "taylor-reid.txt", HERE / "taylor-reid.docx")
    logo(HERE / "northwind-logo.png")
    print("made jordan-patel.pdf, taylor-reid.docx, northwind-logo.png")
